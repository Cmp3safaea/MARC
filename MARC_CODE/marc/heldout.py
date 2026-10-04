"""Structural-core evaluation on an institutional grade record (Section 4).

Applies the frozen R1-R2 configuration, without retuning, to a grade record
with one row per project, and evaluates it against actual grade changes: a
project is reference-positive when the moderator grade or the final grade
differs from the agreed grade.

The grade record may be a .csv or an .xlsx file (xlsx needs openpyxl). The
header row must contain these columns (case-insensitive):

    Student ID
    Supervisor Grade        (or: First Marker Grade)
    Second Marker Grade     (or: Assessor Grade)
    Agreed Grade
    Moderator Grade
    Final Grade

Usage:
    python -m marc.heldout GRADE_RECORD [--sheet NAME] [--label TEXT]
        [--out ranked.csv] [--r2-undefined on|off]

--r2-undefined on   R2 = max(b, u), the Section 2 definition (default).
--r2-undefined off  genuine crossings only. Table 5 as published was
                    produced with this setting.
"""
from __future__ import annotations

import argparse
import csv
import statistics
import sys
from typing import Dict, List, Optional

from .evaluate import confusion, _f, _ci
from .pipeline import Config
from .rules import Project, compute_discrepancy, apply_structural_rules

_ADMIN = {"", "-", "NS", "NEC", "ZERO", "N/A", "NA", "SAI"}
_COLUMNS = {
    "sid": ("student id",),
    "fm": ("supervisor grade", "first marker grade"),
    "as": ("second marker grade", "assessor grade"),
    "agreed": ("agreed grade",),
    "moderator": ("moderator grade",),
    "final": ("final grade",),
}


def _norm(token) -> str:
    return "" if token is None else str(token).strip().upper()


def _table(path: str, sheet: Optional[str]) -> List[List]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError as exc:
            raise SystemExit("reading .xlsx needs openpyxl: "
                             "pip install openpyxl") from exc
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.reader(fh)]


def load_grade_record(path: str, sheet: Optional[str] = None) -> List[dict]:
    table = _table(path, sheet)
    header_at, index = None, {}
    for i, row in enumerate(table[:5]):
        names = [_norm(v).lower() for v in row]
        if "student id" in names:
            header_at = i
            for key, options in _COLUMNS.items():
                index[key] = next((names.index(o) for o in options
                                   if o in names), None)
            break
    if header_at is None:
        raise SystemExit("no header row with a 'Student ID' column found")
    missing = [k for k in ("fm", "as", "agreed") if index.get(k) is None]
    if missing:
        raise SystemExit(f"grade record is missing columns: {missing}")
    rows = []
    for row in table[header_at + 1:]:
        def cell(key):
            j = index.get(key)
            return row[j] if j is not None and j < len(row) else None
        if not _norm(cell("sid")):
            continue
        rows.append({k: cell(k) for k in _COLUMNS})
    return rows


def grade_changed(row: dict) -> bool:
    """Moderator or final grade differs from the agreed grade."""
    agreed = _norm(row["agreed"])
    if agreed in _ADMIN:
        return False
    return any(_norm(row[k]) not in _ADMIN and _norm(row[k]) != agreed
               for k in ("moderator", "final"))


def analyse(rows: List[dict], cfg: Optional[Config] = None,
            r2_include_undefined: bool = True) -> Dict:
    cfg = cfg or Config()
    recs = []
    for row in rows:
        p = Project(student_id=_norm(row["sid"]),
                    fm_token=_norm(row["fm"]), as_token=_norm(row["as"]))
        d = compute_discrepancy(p)
        sf = apply_structural_rules(p, d, cfg.tau_gap, cfg.tau_chars,
                                    r2_include_undefined)
        score = sf.r1 * cfg.weights["r1"] + sf.r2 * cfg.weights["r2"]
        recs.append({
            "student_id": p.student_id,
            "fm_token": p.fm_token, "as_token": p.as_token,
            "agreed": _norm(row["agreed"]),
            "moderator": _norm(row["moderator"]),
            "final": _norm(row["final"]),
            "fm_value": d.fm.value, "as_value": d.asr.value,
            "gap": d.gap, "boundary_crossed": int(d.boundary_crossed),
            "band_undefined": int(d.band_undefined),
            "R1": int(sf.r1), "R2": int(sf.r2), "score": score,
            "grade_changed": int(grade_changed(row))})
    n = len(recs)
    gaps = [r["gap"] for r in recs if r["gap"] is not None]
    crossings = sum(r["boundary_crossed"] for r in recs)
    pos = [bool(r["grade_changed"]) for r in recs]
    selected = confusion([r["score"] >= cfg.cutoff_selected for r in recs], pos)
    burden = selected["flagged"] / n if n else float("nan")
    res = {
        "n": n, "parseable_pairs": len(gaps),
        "unparseable_pairs": n - len(gaps),
        "gap_mean": statistics.fmean(gaps) if gaps else None,
        "gap_median": statistics.median(gaps) if gaps else None,
        "band_crossings": crossings,
        "crossing_rate": crossings / len(gaps) if gaps else None,
        "grade_changes": sum(pos), "base_rate": sum(pos) / n if n else None,
        "selected": selected, "burden": burden,
        "recall_gain_over_chance": (selected["recall"] / burden
                                    if burden else float("nan")),
        "r2_include_undefined": r2_include_undefined,
    }
    return {"summary": res, "records": recs}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        prog="marc.heldout",
        description="Structural core (R1, R2) against actual grade changes.")
    ap.add_argument("grade_record", help=".csv or .xlsx grade record")
    ap.add_argument("--sheet", default=None, help="worksheet name (xlsx)")
    ap.add_argument("--label", default="cohort")
    ap.add_argument("--out", default=None, help="per-project CSV output")
    ap.add_argument("--r2-undefined", choices=("on", "off"), default="on")
    args = ap.parse_args(argv)

    cfg = Config()
    result = analyse(load_grade_record(args.grade_record, args.sheet), cfg,
                     r2_include_undefined=args.r2_undefined == "on")
    s, c = result["summary"], result["summary"]["selected"]
    print(f"Structural core (R1, R2), frozen thresholds: gap >= "
          f"{cfg.tau_gap:g} pp, S >= {cfg.cutoff_selected:g}")
    print(f"R2 undefined-band component: {args.r2_undefined}")
    print(f"\n{args.label}")
    print(f"  projects                    {s['n']}")
    print(f"  pairs with parseable marks  {s['parseable_pairs']} "
          f"({s['unparseable_pairs']} unparseable)")
    print(f"  gap mean / median           {_f(s['gap_mean'], 1)} / "
          f"{_f(s['gap_median'], 1)}")
    print(f"  band-crossing rate          {s['band_crossings']}/"
          f"{s['parseable_pairs']} = {_f(s['crossing_rate'])}")
    print(f"  grade changes (base rate)   {s['grade_changes']} "
          f"({_f(s['base_rate'])})")
    print(f"  flagged (burden)            {c['flagged']} ({_f(s['burden'])})")
    print(f"  precision                   {_f(c['precision'])} "
          f"{_ci(c['precision_ci'])}")
    print(f"  recall                      {_f(c['recall'])} "
          f"{_ci(c['recall_ci'])}")
    print(f"  F1                          {_f(c['F1'])}")
    print(f"  TP {c['TP']}  FP {c['FP']}  FN {c['FN']}  TN {c['TN']}")
    print(f"  recall gain over chance     "
          f"{_f(s['recall_gain_over_chance'], 2)}x")
    if args.out:
        recs = sorted(result["records"], key=lambda r: -r["score"])
        with open(args.out, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(recs[0].keys()))
            w.writeheader()
            w.writerows(recs)
        print(f"\nPer-project results written to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
